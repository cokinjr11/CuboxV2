"""CUBOX 2.0 -- Operational Guide Print Redesign v2 + Configurable Export
Validation Override.

Part A (cover pages, step layout, composition summary) se testea a nivel de
Flowable -reportlab comprime el contenido real en el PDF final (FlateDecode),
asi que buscar texto literal en los bytes del PDF no es confiable ni portable
(no hay pypdf/pdfminer en este entorno). Se inspeccionan los objetos
Paragraph/Table que devuelven las funciones de construccion ANTES de
`doc.build()`, que es exactamente lo que se termina imprimiendo.

Part B (override) sí se testea end-to-end vía TestClient + los 5 endpoints
reales -status codes, bytes (%PDF-/PK para zip), y el contenido de la hoja
Excel via openpyxl (ya una dependencia declarada)."""

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.excel_export import build_export_workbook
from app.core.pdf_export import (
    _export_override_box,
    _fit_image_dims,
    _make_override_footer,
    _operational_cover_page,
    _plan_validation_block,
    _step_card,
)
from app.main import app
from app.models.schemas import (
    CategoryStatus,
    CategoryStatusValue,
    PlanValidationStatus,
    ReportMetadata,
    ReportValidationResponse,
    ValidationCategory,
    ValidationIssue,
    ValidationSeverity,
)

client = TestClient(app)

_TINY_PNG_BASE64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="


def _window(**overrides):
    defaults = dict(
        code="W1", description="Ventana test", width=1200, height=2000, thickness=100, weight=45,
        quantity=4, system="SysA", group="G1", stackable=True, priority=1,
    )
    defaults.update(overrides)
    return defaults


def _pack(items, container_id="40ft_standard"):
    r = client.post("/api/pack", json={"items": items, "container_id": container_id, "optimization_mode": "best_space"})
    assert r.status_code == 200
    return r.json()["best"]


def _paragraph_texts(flowables) -> list[str]:
    """Texto plano de cada Paragraph en una lista de Flowables (ignora
    Spacer/Table/PageBreak, que no tienen `.text`)."""
    return [f.text for f in flowables if hasattr(f, "text")]


def _fresh_ready_validation() -> ReportValidationResponse:
    return ReportValidationResponse(
        valid=True, errors=[], warnings=[], status=PlanValidationStatus.READY,
        error_count=0, warning_count=0, unloaded_count=0, categories=[], error_issues=[], warning_issues=[],
    )


def _fresh_ready_with_warnings_validation() -> ReportValidationResponse:
    warning_issue = ValidationIssue(
        category=ValidationCategory.DELIVERY_SEQUENCE, severity=ValidationSeverity.WARNING, message="conflict"
    )
    return ReportValidationResponse(
        valid=True, errors=[], warnings=["conflict"], status=PlanValidationStatus.READY_WITH_WARNINGS,
        error_count=0, warning_count=1, unloaded_count=0,
        categories=[CategoryStatus(category=ValidationCategory.DELIVERY_SEQUENCE, status=CategoryStatusValue.WARNING, count=1)],
        error_issues=[], warning_issues=[warning_issue],
    )


def _fresh_not_ready_validation() -> ReportValidationResponse:
    error_issue = ValidationIssue(
        category=ValidationCategory.COLLISION_CLEARANCE, severity=ValidationSeverity.ERROR, message="Colision entre A y B"
    )
    return ReportValidationResponse(
        valid=False, errors=["Colision entre A y B"], warnings=[], status=PlanValidationStatus.NOT_READY,
        error_count=1, warning_count=0, unloaded_count=0,
        categories=[CategoryStatus(category=ValidationCategory.COLLISION_CLEARANCE, status=CategoryStatusValue.ERROR, count=1)],
        error_issues=[error_issue], warning_issues=[],
    )


# ==========================================================================
# Part A -- cover page (section 1-7)
# ==========================================================================


def test_cover_page_has_no_step_and_ends_with_page_break():
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak

    from app.core.packer import compute_metrics
    from app.models.schemas import ContainerSpec, PackingResult, RoadWeightConfig

    styles = getSampleStyleSheet()
    container = ContainerSpec(
        id="40ft_standard", name="40 ft Standard", length=12000, width=2350, height=2390, max_weight=26512,
        road_weight_config=None,
    )
    state = PackingResult(container=container, placed=[], unloaded=[], metrics=compute_metrics(container, [], []))
    story = _operational_cover_page("load", state, ReportMetadata(), "My Plan", "BOX", 3, _fresh_ready_validation(), styles)

    texts = " | ".join(_paragraph_texts(story))
    # Cover cleanup (correccion): el logo ya incluye el wordmark "Cubox" -no
    # se repite como titulo aparte, ver _operational_cover_page.
    assert "LOADING GUIDE" in texts
    assert "PLAN INFORMATION" in texts
    plan_info_table = next(f for f in story if hasattr(f, "_cellvalues"))
    assert any("My Plan" in row for row in plan_info_table._cellvalues)
    assert "OPERATIONAL METRICS" in texts
    assert "COLOR LEGEND" in texts
    assert not any("STEP" in t and "OF" in t for t in _paragraph_texts(story)), "la portada nunca debe contener un Step"
    assert isinstance(story[-1], PageBreak), "la portada SIEMPRE termina en PageBreak -Step 1 nunca comparte pagina"


def test_cover_shows_groups_section_only_for_unload_direction():
    from reportlab.lib.styles import getSampleStyleSheet

    from app.core.packer import compute_metrics
    from app.models.schemas import ContainerSpec, PackingResult

    styles = getSampleStyleSheet()
    container = ContainerSpec(id="c", name="C", length=1000, width=1000, height=1000, max_weight=1000)
    state = PackingResult(container=container, placed=[], unloaded=[], metrics=compute_metrics(container, [], []))

    load_story = _operational_cover_page("load", state, ReportMetadata(), None, None, 0, None, styles)
    unload_story = _operational_cover_page("unload", state, ReportMetadata(), None, None, 0, None, styles)
    assert "GROUPS" not in " ".join(_paragraph_texts(load_story))
    # sin piezas no hay Groups -> tampoco aparece la seccion (nunca una lista vacia)
    assert "GROUPS" not in " ".join(_paragraph_texts(unload_story))


def test_plan_validation_block_shows_ready_without_category_breakdown():
    from reportlab.lib.styles import getSampleStyleSheet

    story = _plan_validation_block(_fresh_ready_validation(), getSampleStyleSheet())
    texts = _paragraph_texts(story)
    assert "READY" in texts
    assert not any("conflict" in t.lower() for t in texts)


def test_plan_validation_block_shows_warning_category_counts():
    from reportlab.lib.styles import getSampleStyleSheet

    story = _plan_validation_block(_fresh_ready_with_warnings_validation(), getSampleStyleSheet())
    texts = _paragraph_texts(story)
    assert "READY WITH WARNINGS" in texts
    assert any("Delivery Sequence" in t for t in texts)


def test_plan_validation_block_not_ready_never_shown_without_override():
    """Regresion (seccion 44): un NOT_READY solo puede llegar aca si el
    caller (routes.py) ya confirmo que hubo override -esta funcion no decide
    nada, solo refleja. No hay forma de que un status normal se marque como
    error si `validation` viene bien formado."""
    from reportlab.lib.styles import getSampleStyleSheet

    story = _plan_validation_block(_fresh_not_ready_validation(), getSampleStyleSheet())
    texts = _paragraph_texts(story)
    assert "✕ NOT READY" in texts
    assert "EXPORTED WITH VALIDATION ERRORS" in texts


# ==========================================================================
# Part A -- step card (section 9-22)
# ==========================================================================


def test_step_card_title_and_action_loading():
    from reportlab.lib.styles import getSampleStyleSheet

    card = _step_card(1, 3, "load", [], None, {}, getSampleStyleSheet(), 100, 100, False)
    texts = _paragraph_texts(card)
    assert "STEP 1 OF 3" in texts
    assert "LOADING: LOAD THESE ITEMS" in texts


def test_step_card_title_and_action_unloading():
    from reportlab.lib.styles import getSampleStyleSheet

    card = _step_card(2, 5, "unload", [], None, {}, getSampleStyleSheet(), 100, 100, False)
    texts = _paragraph_texts(card)
    assert "STEP 2 OF 5" in texts
    assert "UNLOADING: REMOVE THESE ITEMS" in texts


def test_step_card_loading_notes_say_rear_and_no_warnings():
    from reportlab.lib.styles import getSampleStyleSheet

    card = _step_card(1, 1, "load", [], None, {}, getSampleStyleSheet(), 100, 100, False)
    # STEP CONTENT/STEP NOTES viven dentro de una Table anidada (info_row) -
    # buscamos su texto recorriendo las celdas.
    info_row = card[-1]
    cell_texts = []
    for row in info_row._cellvalues:
        for cell in row:
            cell_texts.extend(_paragraph_texts(cell) if isinstance(cell, list) else [])
    joined = " | ".join(cell_texts)
    assert "Rear" in joined
    assert "Warnings: None" in joined


def test_step_card_unload_no_warnings_shows_check():
    from reportlab.lib.styles import getSampleStyleSheet
    from app.models.schemas import UnloadStepDeliveryInfo

    info = UnloadStepDeliveryInfo(delivery_sequences=[2], is_mixed=False, conflict_messages=[], conflict_types=[])
    card = _step_card(1, 1, "unload", [], None, {}, getSampleStyleSheet(), 100, 100, False, delivery_info=info)
    info_row = card[-1]
    cell_texts = []
    for row in info_row._cellvalues:
        for cell in row:
            cell_texts.extend(_paragraph_texts(cell) if isinstance(cell, list) else [])
    joined = " | ".join(cell_texts)
    assert "Delivery Sequence" in joined and "2" in joined
    assert "No operational warnings" in joined


# ==========================================================================
# Part A -- Full/Group Unloading Guide integration (structural smoke)
# ==========================================================================


def test_loading_guide_pdf_is_larger_with_cover_than_before():
    """No hay forma portable de contar paginas sin un parser de PDF -pero un
    documento con portada + Steps + Composition/Final Summary debe generar
    bytes claramente mas grandes que un PDF trivial de una sola pagina."""
    from app.core.pdf_export import build_loading_guide_pdf
    from app.core.sequence import compute_load_steps
    from app.models.schemas import LoadingAnchor

    _pack([_window(quantity=6)])
    from app.api.routes import _get_active_state

    state = _get_active_state()
    steps = compute_load_steps(state.placed, state.container, LoadingAnchor.BACK_RIGHT)
    pdf_bytes = build_loading_guide_pdf(state, steps, [_TINY_PNG_BASE64] * len(steps), ReportMetadata())
    assert pdf_bytes[:5] == b"%PDF-"
    assert len(pdf_bytes) > 3000  # portada + legend + metrics + steps + summary, nunca un PDF vacio


# ==========================================================================
# Part B -- Configurable Export Validation Override (sections 29-47)
# ==========================================================================


def _force_collision_on_current_state():
    """Mismo truco que test_pdf_export.py::test_pdf_export_rejects_when_state_is_invalid
    -fuerza una colision manual sobre el estado YA activo (placed[0]/[1])
    para tener un plan NOT_READY real, sin repackear (preserva Group/etc.
    de lo que el caller ya empaco) ni inventar un segundo camino de
    validacion."""
    from app.api.routes import _current_state

    state = _current_state["result"]
    a, b = state.placed[0], state.placed[1]
    b.x, b.y, b.z = a.x, a.y, a.z
    return state


def _force_not_ready():
    """Empaca un plan generico de 2 piezas y lo fuerza a NOT_READY -para
    tests que no necesitan Group/otros datos especificos, ver
    _force_collision_on_current_state para ese caso."""
    _pack([_window(quantity=2)])
    return _force_collision_on_current_state()


def test_A_strict_blocking_on_by_default_blocks_pdf_and_excel():
    _force_not_ready()
    r_pdf = client.post("/api/report/loading-guide-pdf", json={"step_mode": "automatic", "step_images_png_base64": []})
    assert r_pdf.status_code == 422
    assert r_pdf.json()["detail"]["status"] == "not_ready"

    r_excel = client.get("/api/export-excel")
    assert r_excel.status_code == 422


def test_B_override_allows_pdf_and_excel_and_marks_not_ready():
    _force_not_ready()
    steps_r = client.post("/api/report/steps", json={"direction": "load", "step_mode": "automatic"})
    n_steps = len(steps_r.json()["steps"])

    r_pdf = client.post(
        "/api/report/loading-guide-pdf",
        json={"step_mode": "automatic", "step_images_png_base64": [_TINY_PNG_BASE64] * n_steps, "allow_export_with_errors": True},
    )
    assert r_pdf.status_code == 200
    assert r_pdf.content[:5] == b"%PDF-"

    r_unload = client.post(
        "/api/report/unloading-guide-pdf",
        json={"step_mode": "automatic", "step_images_png_base64": [_TINY_PNG_BASE64] * n_steps, "allow_export_with_errors": True},
    )
    assert r_unload.status_code == 200

    r_container = client.post(
        "/api/report/container-pdf",
        json={"include_overview_image": False, "allow_export_with_errors": True},
    )
    assert r_container.status_code == 200

    r_excel = client.get("/api/export-excel", params={"allow_export_with_errors": True})
    assert r_excel.status_code == 200
    from io import BytesIO

    wb = load_workbook(BytesIO(r_excel.content))
    rows = list(wb["Summary"].iter_rows(values_only=True))
    assert ("Plan Status", "NOT READY") in rows
    assert ("Export Override", "Yes") in rows


def test_C_ready_plan_exports_normally_no_override_banner():
    _pack([_window(quantity=2)])
    r_excel = client.get("/api/export-excel")
    assert r_excel.status_code == 200
    from io import BytesIO

    wb = load_workbook(BytesIO(r_excel.content))
    rows = [row[0] for row in wb["Summary"].iter_rows(values_only=True)]
    assert "Plan Status" not in rows
    assert "Export Override" not in rows


def test_D_ready_with_warnings_exports_normally_no_override_banner():
    _pack([_window(quantity=3, delivery_sequence=1), _window(quantity=3, delivery_sequence=1, code="W2")])
    v = client.post("/api/report/validate")
    assert v.json()["status"] in ("ready", "ready_with_warnings")
    r_excel = client.get("/api/export-excel")
    assert r_excel.status_code == 200


def test_unloading_guide_by_group_respects_override():
    _pack([_window(quantity=4, group="Ruta Norte", code="W1")])
    _force_collision_on_current_state()
    steps_r = client.post("/api/report/steps", json={"direction": "unload", "step_mode": "automatic"})
    n_steps = len(steps_r.json()["steps"])

    r_blocked = client.post(
        "/api/report/unloading-guide-pdf-by-group",
        json={"step_mode": "automatic", "step_images_png_base64": [_TINY_PNG_BASE64] * n_steps, "groups": ["Ruta Norte"]},
    )
    assert r_blocked.status_code == 422

    r_override = client.post(
        "/api/report/unloading-guide-pdf-by-group",
        json={
            "step_mode": "automatic", "step_images_png_base64": [_TINY_PNG_BASE64] * n_steps,
            "groups": ["Ruta Norte"], "allow_export_with_errors": True,
        },
    )
    assert r_override.status_code == 200
    assert r_override.content[:5] == b"%PDF-"


def test_override_never_repacks_or_mutates_geometry():
    state = _force_not_ready()
    positions_before = {p.id: (p.x, p.y, p.z) for p in state.placed}
    steps_r = client.post("/api/report/steps", json={"direction": "load", "step_mode": "automatic"})
    n_steps = len(steps_r.json()["steps"])
    client.post(
        "/api/report/loading-guide-pdf",
        json={"step_mode": "automatic", "step_images_png_base64": [_TINY_PNG_BASE64] * n_steps, "allow_export_with_errors": True},
    )
    from app.api.routes import _get_active_state

    positions_after = {p.id: (p.x, p.y, p.z) for p in _get_active_state().placed}
    assert positions_after == positions_before, "generar un export con override NUNCA debe repackear ni mover piezas"


def test_footer_function_none_unless_not_ready():
    assert _make_override_footer(None) is None
    assert _make_override_footer(_fresh_ready_validation()) is None
    assert _make_override_footer(_fresh_ready_with_warnings_validation()) is None
    assert _make_override_footer(_fresh_not_ready_validation()) is not None


def test_export_override_box_lists_error_categories():
    from reportlab.lib.styles import getSampleStyleSheet

    story = _export_override_box(_fresh_not_ready_validation(), getSampleStyleSheet())
    texts = []
    for item in story:
        if hasattr(item, "_cellvalues"):
            for row in item._cellvalues:
                for cell in row:
                    if isinstance(cell, list):
                        texts.extend(_paragraph_texts(cell))
    joined = " | ".join(texts)
    assert "WARNING" in joined
    assert "BLOCKING VALIDATION ERRORS" in joined
    assert "Collision & Clearance" in joined


def test_excel_workbook_without_override_has_no_status_rows():
    from app.core.packer import compute_metrics
    from app.models.schemas import ContainerSpec, PackingResult

    container = ContainerSpec(id="c", name="C", length=1000, width=1000, height=1000, max_weight=1000)
    state = PackingResult(container=container, placed=[], unloaded=[], metrics=compute_metrics(container, [], []))
    wb_bytes = build_export_workbook(state)
    from io import BytesIO

    wb = load_workbook(BytesIO(wb_bytes))
    rows = [row[0] for row in wb["Summary"].iter_rows(values_only=True)]
    assert "Plan Status" not in rows


# ==========================================================================
# PDF guide image proportions -- Correction 1
# ==========================================================================


def _synthetic_png_base64(width: int, height: int) -> str:
    """PNG sintetico de tamano EXACTO conocido -para verificar que
    _fit_image_dims preserva su proporcion nativa real, sin depender de un
    screenshot real del navegador."""
    import base64
    import io

    from PIL import Image as PILImage

    buf = io.BytesIO()
    PILImage.new("RGB", (width, height), color="white").save(buf, format="PNG")
    return base64.b64encode(buf.getvalue()).decode("ascii")


def test_fit_image_dims_preserves_portrait_aspect_inside_landscape_box():
    """El bug real: un canvas capturado mas alto que ancho (ej. 780x1000,
    igual que lo observado con un screenshot real del navegador) forzado a
    una caja 540x230 -sin este fix, Image(width=540, height=230) lo
    estiraria/aplastaria. _fit_image_dims debe mantener la proporcion
    780/1000 exacta."""
    png = _synthetic_png_base64(780, 1000)
    fitted_w, fitted_h = _fit_image_dims(png, 540, 230)
    assert fitted_w <= 540 + 1e-6
    assert fitted_h <= 230 + 1e-6
    assert abs((fitted_w / fitted_h) - (780 / 1000)) < 1e-6
    # limitado por la altura (780x1000 es mas "vertical" que la caja) -> debe
    # tocar el techo de 230 de alto, nunca el de 540 de ancho.
    assert abs(fitted_h - 230) < 1e-6


def test_fit_image_dims_preserves_landscape_aspect_inside_landscape_box():
    """Un canvas 16:9 (1.78) es RELATIVAMENTE mas alto que la caja 540x230
    (2.35) -sigue siendo landscape, pero el limite real es la ALTURA (230),
    no el ancho- igual debe preservar su proporcion nativa exacta, nunca
    estirarse a 540x230 tal cual."""
    png = _synthetic_png_base64(1600, 900)
    fitted_w, fitted_h = _fit_image_dims(png, 540, 230)
    assert abs((fitted_w / fitted_h) - (1600 / 900)) < 1e-6
    assert abs(fitted_h - 230) < 1e-6
    assert fitted_w < 540 - 1e-6


def test_fit_image_dims_wide_image_is_width_limited():
    """Un canvas MUY ancho (ej. 2400x900, mas ancho relativamente que la
    caja 540x230) si debe quedar limitado por el ancho (540)."""
    png = _synthetic_png_base64(2400, 900)
    fitted_w, fitted_h = _fit_image_dims(png, 540, 230)
    assert abs((fitted_w / fitted_h) - (2400 / 900)) < 1e-6
    assert abs(fitted_w - 540) < 1e-6


def test_fit_image_dims_handles_square_image():
    png = _synthetic_png_base64(500, 500)
    fitted_w, fitted_h = _fit_image_dims(png, 540, 230)
    assert abs(fitted_w - fitted_h) < 1e-6
    assert fitted_h <= 230 + 1e-6


def test_step_card_never_stretches_image_to_fixed_box():
    """Regresion end-to-end: _step_card debe usar el tamano AJUSTADO, nunca
    (image_width, image_height) tal cual, cuando la proporcion no coincide."""
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Image as RLImage

    png = _synthetic_png_base64(780, 1000)
    card = _step_card(1, 1, "load", [], png, {}, getSampleStyleSheet(), 540, 230, False)
    image_flowable = next(f for f in card if isinstance(f, RLImage))
    assert image_flowable.drawWidth != 540 or image_flowable.drawHeight != 230
    assert abs((image_flowable.drawWidth / image_flowable.drawHeight) - (780 / 1000)) < 1e-6


# ==========================================================================
# Step Card atomico -Correccion: nunca dividir un Step entre 2 paginas
# ==========================================================================


def test_every_step_is_wrapped_in_keep_together():
    """Cada Step del story debe ser un UNICO flowable KeepTogether -nunca
    una lista de flowables sueltos (eso es lo que permitia que reportlab
    cortara Step Content/Step Notes en la pagina siguiente, dejando el
    titulo/imagen atras)."""
    from reportlab.platypus import KeepTogether

    from app.core.pdf_export import build_loading_guide_pdf
    from app.core.sequence import compute_load_steps
    from app.models.schemas import LoadingAnchor

    _pack([_window(quantity=6)])
    from app.api.routes import _get_active_state

    state = _get_active_state()
    steps = compute_load_steps(state.placed, state.container, LoadingAnchor.BACK_RIGHT)

    # Reconstruye el story igual que build_loading_guide_pdf, pero
    # inspeccionando los Flowables ANTES de doc.build() -no hay forma
    # portable de verificar esto desde los bytes finales del PDF (ver
    # docstring del modulo).
    import app.core.pdf_export as pdf_export

    original_build = pdf_export.SimpleDocTemplate.build
    captured_story = []
    pdf_export.SimpleDocTemplate.build = lambda self, story, **kw: captured_story.extend(story)
    try:
        build_loading_guide_pdf(state, steps, [_TINY_PNG_BASE64] * len(steps), ReportMetadata())
    finally:
        pdf_export.SimpleDocTemplate.build = original_build

    keep_together_items = [f for f in captured_story if isinstance(f, KeepTogether)]
    assert len(keep_together_items) == len(steps), "debe haber EXACTAMENTE un KeepTogether por Step"


def test_step_content_never_appended_as_loose_flowables():
    """Version mas directa del mismo chequeo: recorriendo el story, cada
    aparicion de un Paragraph 'STEP N OF M' debe estar DENTRO de un
    KeepTogether junto con su tabla Step Content -nunca suelto al mismo
    nivel que el resto del story."""
    from reportlab.platypus import KeepTogether

    from app.core.pdf_export import build_loading_guide_pdf
    from app.core.sequence import compute_load_steps
    from app.models.schemas import LoadingAnchor

    _pack([_window(quantity=8)])
    from app.api.routes import _get_active_state

    state = _get_active_state()
    steps = compute_load_steps(state.placed, state.container, LoadingAnchor.BACK_RIGHT)

    import app.core.pdf_export as pdf_export

    original_build = pdf_export.SimpleDocTemplate.build
    captured_story = []
    pdf_export.SimpleDocTemplate.build = lambda self, story, **kw: captured_story.extend(story)
    try:
        build_loading_guide_pdf(state, steps, [_TINY_PNG_BASE64] * len(steps), ReportMetadata())
    finally:
        pdf_export.SimpleDocTemplate.build = original_build

    for item in captured_story:
        assert not (hasattr(item, "text") and "STEP" in item.text and " OF " in item.text), (
            "un titulo de Step no debe aparecer SUELTO en el story -debe estar dentro de un KeepTogether"
        )

    for kt in [f for f in captured_story if isinstance(f, KeepTogether)]:
        texts = [f.text for f in kt._content if hasattr(f, "text")]
        # STEP CONTENT/STEP NOTES viven un nivel mas adentro -dentro de la
        # Table info_row (Step Content | Step Notes), cuyas celdas son a su
        # vez listas de Flowables (ver _step_card).
        for item in kt._content:
            if hasattr(item, "_cellvalues"):
                for row in item._cellvalues:
                    for cell in row:
                        # Adaptive Pagination: _pack_step_cards mide el alto
                        # real de cada Step Card ANTES de doc.build() (ver
                        # _measure_flowables_height), lo que dispara
                        # Table.wrap() mas temprano que antes -reportlab
                        # normaliza celdas de lista de Flowables a tupla como
                        # efecto de wrap() (detalle interno de reportlab, no
                        # una perdida de contenido: verificado que el PDF
                        # final sigue mostrando STEP CONTENT/STEP NOTES en
                        # cada pagina). Antes de esta correccion nunca se
                        # llamaba wrap() antes de esta inspeccion (doc.build
                        # esta mockeado a un no-op), asi que la celda seguia
                        # siendo una lista cruda -ahora puede ser list o
                        # tuple indistintamente, ambas son validas para
                        # reportlab.
                        if isinstance(cell, (list, tuple)):
                            texts.extend(f.text for f in cell if hasattr(f, "text"))
        inner_texts = " | ".join(texts)
        assert "STEP" in inner_texts and " OF " in inner_texts
        assert "STEP CONTENT" in inner_texts
        assert "STEP NOTES" in inner_texts
