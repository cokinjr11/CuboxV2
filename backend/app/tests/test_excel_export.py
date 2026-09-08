"""Exportacion a Excel del resultado de cubicaje (core/excel_export.py).
Cubre especificamente el gating de Boxes Inside (Fase 6.2): la columna solo
debe aparecer para planes Palletized Load, nunca para BOX/PANEL/CUSTOM."""

import io

from openpyxl import load_workbook

from app.core.excel_export import _is_palletized_plan, build_export_workbook
from app.core.packer import compute_metrics, pack_container
from app.models.containers import build_custom_load_space
from app.models.schemas import Dimensions3D, ItemType, LoadSpaceType, OrientationPolicy, PackingResult, WindowItem


def _headers(sheet):
    return [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]


def _build_result(item_type: ItemType, dimensions: Dimensions3D, weight: float) -> PackingResult:
    space = build_custom_load_space("Test Space", LoadSpaceType.CONTAINER, 12000, 2400, 3500, 50000)
    orientation = OrientationPolicy.UPRIGHT if item_type == ItemType.PALLET else OrientationPolicy.FREE
    item = WindowItem(code="X1", dimensions=dimensions, weight=weight, quantity=2, item_type=item_type, orientation_policy=orientation)
    packed = pack_container([item], space, strategy="highest_priority")
    metrics = compute_metrics(space, packed.placed, packed.unloaded)
    return PackingResult(
        container=space, placed=packed.placed, unloaded=packed.unloaded, metrics=metrics,
        load_sequence=[p.id for p in packed.placed], unload_sequence=[p.id for p in packed.placed],
    )


def test_is_palletized_plan_true_for_pallet_pieces():
    result = _build_result(ItemType.PALLET, Dimensions3D(length=1200, width=1000, height=1650), 780)
    assert _is_palletized_plan(result.placed) is True


def test_is_palletized_plan_false_for_box_pieces():
    result = _build_result(ItemType.BOX, Dimensions3D(length=600, width=400, height=300), 25)
    assert _is_palletized_plan(result.placed) is False


def test_is_palletized_plan_false_for_empty_list():
    assert _is_palletized_plan([]) is False


def test_export_workbook_includes_boxes_inside_for_pallet_plan():
    result = _build_result(ItemType.PALLET, Dimensions3D(length=1200, width=1000, height=1650), 780)
    xlsx_bytes = build_export_workbook(result)
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert "Boxes Inside" in _headers(wb["Packing List"])
    assert "Boxes Inside" in _headers(wb["Unloaded Items"])


def test_export_workbook_excludes_boxes_inside_for_box_plan():
    result = _build_result(ItemType.BOX, Dimensions3D(length=600, width=400, height=300), 25)
    xlsx_bytes = build_export_workbook(result)
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    assert "Boxes Inside" not in _headers(wb["Packing List"])
    assert "Boxes Inside" not in _headers(wb["Unloaded Items"])
